from __future__ import annotations

import sqlite3
from pathlib import Path

from factorias import crear_carton, crear_evento, crear_lote, crear_organizacion

from bingo.dominio.modelos import Comprador
from bingo.persistencia import repo_carton, repo_comprador
from bingo.servicios import servicio_conciliacion

_TEXTOS = {"titulo": "Conciliación"}


def _evento_con_cartones(con: sqlite3.Connection, *, precio_centavos: int = 100):
    org = crear_organizacion(con)
    evento = crear_evento(con, org.id, precio_tabla_centavos=precio_centavos)
    lote = crear_lote(con, evento.id)
    cartones = [crear_carton(con, evento.id, lote.id, semilla=i) for i in range(5)]
    return evento, cartones


def test_calcular_cuenta_los_cinco_estados(con: sqlite3.Connection) -> None:
    evento, cartones = _evento_con_cartones(con)
    estados = ["generado", "impreso", "entregado", "vendido", "anulado"]
    for carton, estado in zip(cartones, estados, strict=True):
        repo_carton.actualizar_estado(con, carton.id, estado)
    conciliacion = servicio_conciliacion.calcular(con, evento.id)
    assert conciliacion.generados == 1
    assert conciliacion.impresos == 1
    assert conciliacion.entregados == 1
    assert conciliacion.vendidos == 1
    assert conciliacion.anulados == 1
    assert conciliacion.total_cartones == 5


def test_recaudacion_teorica_es_entera(con: sqlite3.Connection) -> None:
    evento, cartones = _evento_con_cartones(con, precio_centavos=333)
    for carton in cartones:
        repo_carton.actualizar_estado(con, carton.id, "vendido")
        repo_comprador.crear(con, Comprador(carton_id=carton.id, nombre="X"))
    conciliacion = servicio_conciliacion.calcular(con, evento.id)
    assert conciliacion.recaudacion_teorica_centavos == 333 * 5
    assert isinstance(conciliacion.recaudacion_teorica_centavos, int)


def test_discrepancia_cuando_vendido_sin_comprador(con: sqlite3.Connection) -> None:
    evento, cartones = _evento_con_cartones(con)
    repo_carton.actualizar_estado(con, cartones[0].id, "vendido")  # sin comprador: anómalo
    conciliacion = servicio_conciliacion.calcular(con, evento.id)
    assert conciliacion.discrepancia is True


def test_sin_discrepancia_con_provisionales(con: sqlite3.Connection) -> None:
    """Hallazgo E-C1: comparar contra compradores vivos (provisionales +
    registrados), no solo contra los registrados — si no, toda venta de
    puerta dispararía una falsa alarma permanente."""
    from bingo.dominio.modelos import Comprador

    evento, cartones = _evento_con_cartones(con)
    for carton in cartones:
        repo_carton.actualizar_estado(con, carton.id, "vendido")
        repo_comprador.crear(con, Comprador(carton_id=carton.id, nombre="Puerta", provisional=True))
    conciliacion = servicio_conciliacion.calcular(con, evento.id)
    assert conciliacion.compradores_provisionales == 5
    assert conciliacion.discrepancia is False


def test_declarar_recaudado_real_y_diferencia(con: sqlite3.Connection) -> None:
    evento, cartones = _evento_con_cartones(con, precio_centavos=100)
    for carton in cartones:
        repo_carton.actualizar_estado(con, carton.id, "vendido")
        repo_comprador.crear(con, Comprador(carton_id=carton.id, nombre="X"))
    servicio_conciliacion.declarar_recaudado_real(con, evento.id, 600)
    conciliacion = servicio_conciliacion.calcular(con, evento.id)
    assert conciliacion.recaudado_real_centavos == 600
    assert conciliacion.diferencia_centavos == 600 - 500


def test_generar_reporte_excel_sin_contacto_por_defecto(
    con: sqlite3.Connection, tmp_path: Path
) -> None:
    from bingo.dominio.modelos import Comprador

    evento, cartones = _evento_con_cartones(con)
    repo_carton.actualizar_estado(con, cartones[0].id, "vendido")
    repo_comprador.crear(
        con, Comprador(carton_id=cartones[0].id, nombre="Ana", telefono="0999", cedula="123")
    )
    ruta = servicio_conciliacion.generar_reporte_excel(con, evento.id, tmp_path / "c.xlsx", _TEXTOS)
    assert ruta.exists()

    import openpyxl

    wb = openpyxl.load_workbook(ruta)
    detalle = wb["Detalle"]
    encabezados = [c.value for c in next(detalle.iter_rows(min_row=1, max_row=1))]
    assert "telefono" not in encabezados
    assert "codigo" in encabezados


def test_generar_reporte_excel_incluye_contacto_si_se_pide(
    con: sqlite3.Connection, tmp_path: Path
) -> None:
    from bingo.dominio.modelos import Comprador

    evento, cartones = _evento_con_cartones(con)
    repo_carton.actualizar_estado(con, cartones[0].id, "vendido")
    repo_comprador.crear(con, Comprador(carton_id=cartones[0].id, nombre="Ana", telefono="0999"))
    ruta = servicio_conciliacion.generar_reporte_excel(
        con, evento.id, tmp_path / "c.xlsx", _TEXTOS, incluir_contacto=True
    )
    import openpyxl

    wb = openpyxl.load_workbook(ruta)
    detalle = wb["Detalle"]
    encabezados = [c.value for c in next(detalle.iter_rows(min_row=1, max_row=1))]
    assert "telefono" in encabezados


def test_generar_reporte_pdf(con: sqlite3.Connection, tmp_path: Path) -> None:
    evento, _cartones = _evento_con_cartones(con)
    ruta = servicio_conciliacion.generar_reporte_pdf(con, evento.id, tmp_path / "c.pdf", _TEXTOS)
    assert ruta.exists()
    assert ruta.stat().st_size > 0


def _evento_con_ronda_y_ganador(con: sqlite3.Connection):
    from factorias import crear_ganador, crear_patron, crear_ronda

    evento, cartones = _evento_con_cartones(con)
    repo_carton.actualizar_estado(con, cartones[0].id, "vendido")
    repo_comprador.crear(con, Comprador(carton_id=cartones[0].id, nombre="Ana"))
    patron = crear_patron(con)
    ronda = crear_ronda(con, evento.id, patron.id, orden=1, nombre="Ronda 1")
    crear_ganador(con, ronda.id, cartones[0].id, bola_numero=7, decision="unico")
    return evento, ronda, cartones[0]


def test_generar_reporte_evento_excel_tiene_hoja_de_rondas(
    con: sqlite3.Connection, tmp_path: Path
) -> None:
    evento, _ronda, carton = _evento_con_ronda_y_ganador(con)
    ruta = servicio_conciliacion.generar_reporte_evento_excel(
        con, evento.id, tmp_path / "r.xlsx", _TEXTOS
    )
    import openpyxl

    wb = openpyxl.load_workbook(ruta)
    assert "Rondas" in wb.sheetnames
    filas = list(wb["Rondas"].iter_rows(min_row=2, values_only=True))
    assert any(carton.codigo in fila for fila in filas)


def test_generar_reporte_evento_excel_sin_contacto_por_defecto(
    con: sqlite3.Connection, tmp_path: Path
) -> None:
    evento, _ronda, _carton = _evento_con_ronda_y_ganador(con)
    ruta = servicio_conciliacion.generar_reporte_evento_excel(
        con, evento.id, tmp_path / "r.xlsx", _TEXTOS
    )
    import openpyxl

    wb = openpyxl.load_workbook(ruta)
    encabezados_rondas = [c.value for c in next(wb["Rondas"].iter_rows(min_row=1, max_row=1))]
    assert "telefono" not in encabezados_rondas


def test_generar_reporte_evento_pdf_no_lleva_contacto(
    con: sqlite3.Connection, tmp_path: Path
) -> None:
    evento, _ronda, _carton = _evento_con_ronda_y_ganador(con)
    ruta = servicio_conciliacion.generar_reporte_evento_pdf(
        con, evento.id, tmp_path / "r.pdf", _TEXTOS
    )
    assert ruta.exists()
    assert ruta.stat().st_size > 0


def test_generar_reporte_evento_sin_ganadores_lista_la_ronda_vacia(
    con: sqlite3.Connection, tmp_path: Path
) -> None:
    from factorias import crear_patron, crear_ronda

    evento, _cartones = _evento_con_cartones(con)
    patron = crear_patron(con)
    crear_ronda(con, evento.id, patron.id, orden=1, nombre="Ronda sin ganador")
    ruta = servicio_conciliacion.generar_reporte_evento_pdf(
        con, evento.id, tmp_path / "r.pdf", _TEXTOS
    )
    assert ruta.exists()
