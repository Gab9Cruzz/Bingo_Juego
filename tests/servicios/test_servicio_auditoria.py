import sqlite3
from pathlib import Path

import openpyxl
from factorias import crear_evento, crear_organizacion

from bingo.persistencia import repo_auditoria
from bingo.servicios import servicio_auditoria

_TEXTOS = {
    "hoja": "Auditoría",
    "columna_momento": "Momento",
    "columna_accion": "Acción",
    "columna_detalle": "Detalle",
}


def _preparar(con: sqlite3.Connection):
    org = crear_organizacion(con)
    evento = crear_evento(con, org.id)
    repo_auditoria.registrar(con, evento.id, "sorteo.ronda_iniciada", detalle="1")
    repo_auditoria.registrar(con, evento.id, "venta.anulada", detalle="carton=A-1")
    repo_auditoria.registrar(con, evento.id, "venta.registrada", detalle="carton=A-2")
    return con, evento


def test_listar_sin_filtro_devuelve_todo(con: sqlite3.Connection) -> None:
    con, evento = _preparar(con)
    registros = servicio_auditoria.listar(con, evento.id)
    assert len(registros) == 3


def test_listar_con_prefijo_filtra(con: sqlite3.Connection) -> None:
    con, evento = _preparar(con)
    registros = servicio_auditoria.listar(con, evento.id, prefijo="venta.")
    assert len(registros) == 2
    assert all(r.accion.startswith("venta.") for r in registros)


def test_listar_prefijo_vacio_no_filtra(con: sqlite3.Connection) -> None:
    con, evento = _preparar(con)
    assert len(servicio_auditoria.listar(con, evento.id, prefijo="")) == 3


def test_generar_reporte_excel_respeta_el_mismo_filtro(
    con: sqlite3.Connection, tmp_path: Path
) -> None:
    con, evento = _preparar(con)
    ruta = tmp_path / "auditoria.xlsx"
    servicio_auditoria.generar_reporte_excel(con, evento.id, ruta, _TEXTOS, prefijo="venta.")

    libro = openpyxl.load_workbook(ruta)
    hoja = libro["Auditoría"]
    filas = list(hoja.iter_rows(min_row=2, values_only=True))
    libro.close()
    assert len(filas) == 2
    assert all(fila[1].startswith("venta.") for fila in filas)


def test_generar_reporte_excel_encabezados_traducidos(
    con: sqlite3.Connection, tmp_path: Path
) -> None:
    con, evento = _preparar(con)
    ruta = tmp_path / "auditoria.xlsx"
    servicio_auditoria.generar_reporte_excel(con, evento.id, ruta, _TEXTOS)

    libro = openpyxl.load_workbook(ruta)
    hoja = libro.active
    encabezados = [c.value for c in hoja[1]]
    libro.close()
    assert encabezados == ["Momento", "Acción", "Detalle"]
